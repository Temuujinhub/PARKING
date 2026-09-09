"""Гэрээт машины жагсаалтыг Excel-ээс уншиж бүртгэх.

Бодит файлууд нь «нэг хуудас = нэг түрээслэгч байгууллага» бүтэцтэй, хуудас бүр
өөр өөр гарчигтай, багануудын байрлал ч зөрдөг. Тиймээс тогтмол баганад
найдахгүй — гарчгийн МӨРИЙГ хайж, «Улсын дугаар» гэх мэт нэрээр багануудыг
олно. Нэг мөрөнд хэд хэдэн дугаар байж болно (ж: банкны мөнгө зөөвөрлөх машин).

Гол хэрэглээ:
    rows, warnings = parse_workbook(file_bytes)
    → rows: [{plate, full_name, note, company, sheet}]
"""
import io
import re


def _site_tenant_id(db, site_id):
    """Зогсоолын түрээслэгч — импортоор үүссэн бүртгэл зогсоолынхоо
    түрээслэгчид харьяалагдана (түрээслэгч дамнахгүй)."""
    if not site_id:
        return None
    from ..models import ParkingSite
    return db.query(ParkingSite.tenant_id).filter(ParkingSite.id == site_id).scalar()


# Гарчгийн нүдийг таних түлхүүр үгс (жижиг үсгээр харьцуулна)
PLATE_HEADERS = ("улсын дугаар", "улсын дугаар ", "дугаар", "plate")
NAME_HEADERS = ("эзэмшигч", "нэр", "owner")
NOTE_HEADERS = ("албан тушаал", "тушаал", "position", "тэмдэглэл")
# «Утас», «Утасны дугаар» — «дугаар» гэсэн үгтэй тул улсын дугаарын баганатай
# андуурагдахгүйн тулд тусад нь таньж, PLATE_HEADERS-ээс хасна
PHONE_HEADERS = ("утас", "phone")

# Загвар файлын ЯГ ТОГТСОН гарчиг (Бүртгэлтэй машин → «Excel загвар татах»).
# parse_workbook эдгээр нэрээр баганыг олдог тул хэрэглэгч баганын ДАРААЛЛЫГ
# солисон ч уншигдана; харин нэрийг нь өөрчилбөл олдохгүй.
TEMPLATE_HEADERS = ("Улсын дугаар", "Эзэмшигч", "Утас", "Албан тушаал")

# 4 орон + 3 кирилл (1234УБА) эсвэл дипломат/тусгай: 2 үсэг + 4 орон (ДК1234).
# Хооронд нь зай/зураас байж болно.
PLATE_RE = re.compile(r"^(?:\d{4}[А-ЯӨҮЁ]{3}|[А-ЯӨҮЁ]{2}\d{4})$")


def normalize_plate(raw: str) -> str:
    """«76-49 УБЯ» → «7649УБЯ». Зай, зураас, цэг, latin-ийг кирилл рүү."""
    s = (raw or "").upper().strip()
    s = re.sub(r"[\s\-–—_.,/\\'\"«»()]", "", s)
    # Excel-д латин үсэг холилдож бичигдсэн тохиолдол элбэг
    for lat, cyr in (("A", "А"), ("B", "В"), ("C", "С"), ("E", "Е"), ("H", "Н"),
                     ("K", "К"), ("M", "М"), ("O", "О"), ("P", "Р"), ("T", "Т"),
                     ("X", "Х"), ("Y", "У")):
        s = s.replace(lat, cyr)
    return s


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    return "" if v is None else str(v).strip()


def _find_header(rows: list) -> tuple[int | None, list[int], int | None, int | None, int | None]:
    """Гарчгийн мөрийг олж (индекс, дугаарын баганууд, нэр, тэмдэглэл, утас) буцаана."""
    for i, row in enumerate(rows[:12]):  # гарчиг эхний мөрүүдэд байдаг
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        phone_col = next((j for j, c in enumerate(cells)
                          if any(h in c for h in PHONE_HEADERS)), None)
        plate_cols = [j for j, c in enumerate(cells)
                      if any(h in c for h in PLATE_HEADERS) and c and j != phone_col]
        if not plate_cols:
            continue
        name_col = next((j for j, c in enumerate(cells)
                         if any(h == c or h in c for h in NAME_HEADERS)), None)
        note_col = next((j for j, c in enumerate(cells)
                         if any(h in c for h in NOTE_HEADERS)), None)
        # Гарчгийн мөрөнд байхгүй ч дугаартай зэрэгцээ багана (ж: "банкны мөнгөн
        # зөөврийн машин") байвал түүнийг ч дугаарын багана гэж үзнэ
        for j, c in enumerate(cells):
            if c and j not in plate_cols and "машин" in c and j > max(plate_cols):
                plate_cols.append(j)
        return i, sorted(plate_cols), name_col, note_col, phone_col
    return None, [], None, None, None


def normalize_phone(raw) -> str:
    """«9911-2233», «+976 99112233», 99112233.0 (Excel тоо) → «99112233». 8 оронтой
    биш бол хоосон — буруу утас хадгалснаас хоосон нь дээр."""
    if raw is None:
        return ""
    s = str(raw).strip()
    if s.endswith(".0"):
        s = s[:-2]
    digits = re.sub(r"\D", "", s)
    if digits.startswith("976") and len(digits) == 11:
        digits = digits[3:]
    return digits if len(digits) == 8 else ""


def _sheet_title(rows: list, sheet_name: str, hdr: int = 3) -> str:
    """Гарчгийн мөрөөс ДЭЭШХИ эхний утгатай нүдийг байгууллагын нэр болгоно
    (ж: «"SGS" гадна автомашины зогсоолын бүртгэл» → «SGS»). Гарчиг хамгийн эхний
    мөрөнд байвал (загвар файл) хуудасны нэр = байгууллагын нэр. Өмнө нь гарчгийн
    мөрийг ч шалгадаг байсан тул загвар файлын байгууллага «Улсын дугаар» болдог байв."""
    for row in rows[:min(3, hdr)]:
        for c in row:
            if c and str(c).strip():
                t = str(c).strip()
                t = re.sub(r"\s*гадна\s+автомашины.*$", "", t, flags=re.I).strip()
                t = t.replace('"', "").replace("«", "").replace("»", "").strip()
                if t and not t.startswith("/"):
                    return t[:160]
    return sheet_name.strip()[:160]


def parse_workbook(data: bytes) -> tuple[list[dict], list[str]]:
    """Excel-ийн БҮХ хуудсыг уншиж мөрүүд + анхааруулга буцаана."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out: list[dict] = []
    warnings: list[str] = []
    seen: set[str] = set()

    for ws in wb.worksheets:
        if ws.title.strip().upper() == "ЗААВАР":
            continue  # загвар файлын зааврын хуудас — доторх текст нь гарчиг мэт харагдана
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr, plate_cols, name_col, note_col, phone_col = _find_header(rows)
        if hdr is None:
            warnings.append(f"«{ws.title}»: «Улсын дугаар» багана олдсонгүй — алгаслаа")
            continue
        company = _sheet_title(rows, ws.title, hdr)
        found = 0
        for row in rows[hdr + 1:]:
            name = _cell(row, name_col)
            note = _cell(row, note_col)
            phone = normalize_phone(row[phone_col]) if phone_col is not None and phone_col < len(row) else ""
            for pc in plate_cols:
                plate = normalize_plate(_cell(row, pc))
                if not plate:
                    continue
                if not PLATE_RE.match(plate):
                    # Гарчиг давтагдсан, нийлбэр мөр гэх мэт — чимээгүй алгасна,
                    # гэхдээ дугаар мэт харагдвал анхааруулна
                    if any(ch.isdigit() for ch in plate) and len(plate) >= 5:
                        warnings.append(f"«{company}»: танигдахгүй дугаар «{_cell(row, pc)}»")
                    continue
                if plate in seen:
                    warnings.append(f"«{company}»: {plate} давхардсан — эхнийхийг үлдээв")
                    continue
                seen.add(plate)
                out.append({"plate": plate, "full_name": name[:120], "note": note[:200],
                            "phone": phone, "company": company, "sheet": ws.title.strip()})
                found += 1
        if not found:
            warnings.append(f"«{company}»: нэг ч дугаар олдсонгүй")
    return out, warnings


def import_rows(db, rows: list[dict], site_id: str | None, *,
                contract_type: str = "CONTRACT", valid_days: int = 365,
                monthly_fee: float = 0, deactivate_missing: bool = False,
                default_tenant_id: str | None = None,
                access_scope: str = "site") -> dict:
    """Задалсан мөрүүдийг registered_drivers руу оруулна (идемпотент upsert).

    Түлхүүр = (plate_number, site_id). Байвал шинэчилнэ, байхгүй бол үүсгэнэ —
    файлыг олон удаа импортлож болно (давхардал үүсэхгүй).

    deactivate_missing=True бол тухайн зогсоолын жагсаалтад ОРООГҮЙ хуучин
    бүртгэлүүдийг идэвхгүй болгоно (жагсаалтыг файлаар бүрэн солих горим).
    Устгадаггүй — түүх, буруу импортоос сэргээх боломж хадгалагдана.
    """
    from datetime import datetime, timedelta

    from ..models import RegisteredDriver

    now = datetime.utcnow()
    valid_to = now + timedelta(days=valid_days)

    # Ижил дугаартай ХЭД ХЭДЭН мөр байж болно (хуучин давхардал): хамгийн урт
    # хүчинтэйг нь «гол» болгож шинэчилнэ, бусдыг нь идэвхгүй болгоно — эс бол
    # dict-д сүүлийнх нь л үлдэж, нөгөө хоёр нь хөндөгдөлгүй давхардсаар байв.
    existing: dict[str, RegisteredDriver] = {}
    deduped = 0
    for d in sorted(db.query(RegisteredDriver).filter(RegisteredDriver.site_id == site_id).all(),
                    key=lambda x: (x.is_active, x.valid_to or datetime.min), reverse=True):
        if d.plate_number in existing:
            if d.is_active:
                d.is_active = False
                d.note = f"{d.note + ' | ' if d.note else ''}импорт: давхардал — {existing[d.plate_number].id} үлдээв"[:1000]
                deduped += 1
            continue
        existing[d.plate_number] = d
    created = updated = 0
    for r in rows:
        d = existing.get(r["plate"])
        if d:
            d.full_name = r["full_name"] or d.full_name
            d.phone = r.get("phone") or d.phone
            d.company = r["company"]
            d.note = r["note"]
            d.contract_type = contract_type
            d.access_scope = access_scope
            d.valid_to = valid_to
            d.is_active = True
            updated += 1
        else:
            db.add(RegisteredDriver(
                plate_number=r["plate"], full_name=r["full_name"], company=r["company"],
                phone=r.get("phone") or "",
                note=r["note"], contract_type=contract_type, site_id=site_id,
                access_scope=access_scope,
                tenant_id=_site_tenant_id(db, site_id) or default_tenant_id,
                monthly_fee=monthly_fee, valid_from=now, valid_to=valid_to, is_active=True))
            created += 1

    deactivated = 0
    if deactivate_missing:
        keep = {r["plate"] for r in rows}
        for plate, d in existing.items():
            if plate not in keep and d.is_active:
                d.is_active = False
                deactivated += 1

    db.commit()
    return {"created": created, "updated": updated, "deactivated": deactivated,
            "deduped": deduped, "total": len(rows)}


def build_template() -> bytes:
    """Импортын ЗАГВАР .xlsx — Бүртгэлтэй машин → «Excel загвар татах».

    Хэрэглэгчид өөр өөр баганатай файл ирүүлж, багана зөрж уншигддаг байсан тул
    яг тогтсон жишээ өгнө: «ЗААВАР» хуудас + байгууллага тус бүрийн хуудас
    (хуудасны нэр = байгууллагын нэр, гарчиг 1-р мөрөнд, TEMPLATE_HEADERS).
    parse_workbook энэ файлыг өөрчлөлтгүй уншиж чадах ёстой (тест шалгадаг)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    guide = wb.active
    guide.title = "ЗААВАР"
    lines = [
        ("Бүртгэлтэй машин — Excel импортын загвар", True),
        ("", False),
        ("1. Хуудас бүр = НЭГ байгууллага. Хуудасны нэрийг байгууллагын нэрээр солино "
         "(ж: «Байгууллага 1» → «Монгол Банк»). Байгууллага олон бол хуудас нэмнэ.", False),
        ("2. 1-р мөрийн гарчгийг ӨӨРЧЛӨХГҮЙ: " + " | ".join(TEMPLATE_HEADERS)
         + ". Баганын дараалал чухал биш, харин НЭР нь яг ийм байх ёстой.", False),
        ("3. «Улсын дугаар» — заавал. 4 орон + 3 кирилл үсэг (1234УБА) эсвэл дипломат "
         "2 үсэг + 4 орон (ДК1234). Зай, зураас, латин үсэг байсан ч систем цэвэрлэнэ.", False),
        ("4. «Эзэмшигч», «Утас» (8 орон), «Албан тушаал» — заавал биш, хоосон үлдээж болно.", False),
        ("5. Жишээ мөрүүдийг устгаад өөрийн жагсаалтаа бичнэ. Ижил дугаар давхардвал "
         "нэг л удаа бүртгэгдэнэ.", False),
        ("6. Импорт хийхдээ зогсоол, бүртгэлийн төрөл (гэрээт/сарын/…), давхар "
         "зогсоолд хамрах хүрээг цонхноос сонгоно. Эхлээд «урьдчилан харах» гарна — "
         "тоо, байгууллага зөв бол баталгаажуулна.", False),
        ("Энэ «ЗААВАР» хуудсыг устгах шаардлагагүй — импорт алгасна.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = guide.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold, size=13 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 110

    head_fill = PatternFill("solid", fgColor="DDEBF7")
    samples = {
        "Байгууллага 1": [
            ("1234УБА", "Бат-Эрдэнэ", "99112233", "Жолооч"),
            ("5678УНЕ", "Сарнай", "", "Захирал"),
            ("ДК1234", "", "", "дипломат дугаар мөн болно"),
        ],
        "Байгууллага 2": [
            ("2345УБВ", "Дорж", "88001122", "хуудас бүр тусдаа байгууллага болно"),
        ],
    }
    for title, rows in samples.items():
        ws = wb.create_sheet(title)
        ws.append(list(TEMPLATE_HEADERS))
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate((16, 24, 14, 40), start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
