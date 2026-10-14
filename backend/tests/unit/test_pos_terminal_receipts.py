"""POS терминал (PAX апп) өөрөө НӨАТ баримт гаргадаг — давхар e-Barimt (2026-10-13; 10-14).

ТЕГ тулгалт (прод хуулбар): систем карт/бэлэн/дансаар төлбөр бүрд msgbill-ээр ЯГ НЭГ
баримт гаргадаг; давхардал нь PAX терминал Ontime-ийн (касс …10002990) баримтыг
ӨӨРӨӨ хэвлэдэг ч апп тэр ДДТД-г /pos/confirm, /cash-д илгээдэггүйгээс үүсдэг —
2026-08-24 Хангарьд 63 (39 бэлэн + 24 карт, pos.khangarid). 2026-10-14 эзний эцсийн
шийдвэр: терминал БҮХ зогсоолд хэвлэдэг → ерөнхий анхдагч АСААЛТТАЙ, зогсоолоор унтраана.

Сүлжээ/банк/msgbill огт дуудахгүй — бүх суваг stub.
"""
import asyncio
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import httpx
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from starlette.requests import Request

from app import models as M, session_logic as SL
from app.config import settings
from app.database import Base
from app.routers import compensations_router as C, payments_router as PR, reports_router as R
from app.routers import reports_excel as X, vat_recon as V
from app.services import app_settings as A, financial_jobs as J, payment_rules as PRULES
from app.services import payment_wait, vat_duplicates as VD
from app.services.msgbill import MsgbillAccount
from app.services.receipts import choose_primary

ACC = MsgbillAccount(api_key='synthetic-key', base_url='https://msgbill.invalid')


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    async def blocked_async(*a, **k):
        raise AssertionError('Network forbidden')
    monkeypatch.setattr(httpx.AsyncClient, 'send', blocked_async)
    monkeypatch.setattr(settings, 'qpay_ebarimt', False)
    # msgbill тохируулсан зогсоол (прод шиг): карт/бэлэн → MSGBILL суваг
    monkeypatch.setattr(J.msgbill, 'account_enabled_for', lambda *a: ACC)
    monkeypatch.setattr(J.msgbill, 'api_key_for', lambda *a: ACC)


@pytest.fixture(autouse=True)
def fresh_pos_app_cache():
    from app.services import pos_app as PA
    PA.invalidate_cache()
    yield
    PA.invalidate_cache()


@pytest.fixture
def db():
    engine = create_engine('sqlite://', poolclass=StaticPool, connect_args={'check_same_thread': False})
    Base.metadata.create_all(engine)
    A.invalidate_cache()
    with Session(engine, autoflush=False) as session:
        yield session
    A.invalidate_cache()
    engine.dispose()


def parking(db, code='HANGARID', plate='1234УБА'):
    tariff = M.TariffTemplate(name='Hourly', free_minutes=0, extra_hour_price=1000, daily_cap=50000)
    site = M.ParkingSite(name='Хангарьд', site_code=code, tariff_template=tariff)
    db.add(site); db.flush()
    now = datetime.utcnow()
    s = M.ParkingSession(site_id=site.id, plate_number=plate, entry_time=now - timedelta(minutes=59),
                         status='AWAITING_PAYMENT')
    db.add(s); db.flush(); db.refresh(s)
    payment_wait.begin_wait(db, s, SL.session_fee_info(db, s, at=now), now)
    db.commit()
    return s


def user(db, name='pos.khangarid'):
    u = M.User(username=name, role='SUPER_ADMIN', password_hash='unused')
    db.add(u); db.commit()
    return u


def req(app_version=None, ua=None):
    headers = [(b'x-app-version', app_version.encode())] if app_version else []
    if ua:
        headers.append((b'user-agent', ua.encode()))
    return Request({'type': 'http', 'method': 'POST', 'path': '/', 'headers': headers,
                    'client': ('pos', 1)})


def enable(db, site_id, methods='CARD,CASH', on=True):
    A.set_site_rules(db, A.EBARIMT_KEY, site_id,
                     {'pos_terminal_issues': on, 'pos_terminal_methods': methods}, 'admin')
    db.commit()


def card(db, s, u, txn='bank-1'):
    return asyncio.run(PR.pos_confirm({'session_id': s.id, 'amount': 1000, 'transaction_id': txn},
                                      req(), db, u))


def cash(db, s, u, app_version=None, ua=None):
    return asyncio.run(PR.cash_payment({'session_id': s.id}, db, u, req(app_version, ua)))


def receipts(db, payment_id):
    return db.query(M.VatReceipt).filter_by(payment_id=payment_id).all()


def jobs(db, payment_id):
    return db.query(M.FinancialJob).filter_by(payment_id=payment_id, kind='RECEIPT').all()


# ── 1. Зогсоолд УНТРААСАН — хуучин зан: msgbill баримт + ажил ────────────────
def test_setting_off_keeps_msgbill_receipt_and_job(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id, on=False)                   # терминал нь хэвлэдэггүй зогсоол
    out = card(db, s, u)
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP'                       # /pos/confirm = үргэлж POS апп
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING')
    [j] = jobs(db, p.id)
    assert j.status == 'PENDING' and j.payload['provider'] == 'MSGBILL'
    assert out['receipt']['ebarimt_status'] == 'PENDING'
    # outbox 2–5с — ДДТД хүлээгдэж буйг шударгаар хэвлэнэ («ДДТД: -» биш)
    assert 'ДДТД: хүлээгдэж байна' in out['print_data']['lines']


def test_default_is_on_for_every_site(db):
    """2026-10-14: ерөнхий анхдагч АСААЛТТАЙ — тохиргоо огт бичээгүй зогсоолд ч, зогсоол
    тодорхойгүй төлбөрт ч (ерөнхий утга)."""
    assert A.get_rules(db, A.EBARIMT_KEY)['pos_terminal_issues'] is True
    assert A.pos_terminal_receipt_methods(db, 'any-site') == {'CARD', 'CASH'}
    assert A.pos_terminal_receipt_methods(db, None) == {'CARD', 'CASH'}


def test_default_on_pos_card_and_header_cash_are_external_without_any_setting(db):
    s, u = parking(db), user(db)
    out = card(db, s, u)                                # тохиргоо БИЧЭЭГҮЙ зогсоол
    [r] = receipts(db, out['payment_id'])
    assert (r.provider, r.status) == ('POS-TERMINAL', 'EXTERNAL') and jobs(db, out['payment_id']) == []
    s2 = another_session(db, s.site_id, '4444УБА')
    out2 = cash(db, s2, u, app_version='1.5.3')
    assert receipts(db, out2['payment_id'])[0].status == 'EXTERNAL'
    assert db.query(M.FinancialJob).count() == 0


# ── 2. Асаалттай + POS карт → EXTERNAL, ажилгүй ────────────────────────────
def test_setting_on_pos_card_is_external_without_job(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    out = card(db, s, u)
    p = db.get(M.Payment, out['payment_id'])
    [r] = receipts(db, p.id)
    assert (r.provider, r.status, r.ebarimt_id) == ('POS-TERMINAL', 'EXTERNAL', None)
    assert r.receipt_url == J.TERMINAL_NOTE and r.ddtd_note is None
    assert r.amount == p.amount
    assert jobs(db, p.id) == [] and db.query(M.FinancialJob).count() == 0
    # Хэвлэх: «ДДТД: -» биш, терминалын мөр; төлөв EXTERNAL
    lines = out['print_data']['lines']
    assert 'НӨАТ баримт: POS терминал (Ontime) гаргасан' in lines
    assert not any(ln.startswith('ДДТД') for ln in lines)
    assert out['receipt']['ebarimt_status'] == 'EXTERNAL' and out['ebarimt_id'] is None
    assert out['receipt']['ebarimt_note'] == J.TERMINAL_NOTE


def test_app_sent_ebarimt_id_keeps_existing_external_path(db):
    """Апп терминалын ДДТД-г илгээвэл (урт хугацааны зөв шийдэл) — тэр нь SENT."""
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    body = {'session_id': s.id, 'amount': 1000, 'transaction_id': 'bank-x',
            'ebarimt_id': '029100244106001097210230010045912', 'lottery_code': 'XA 1',
            'ebarimt_provider': 'ONTIME'}
    out = asyncio.run(PR.pos_confirm(body, req('1.4.0'), db, u))
    [r] = receipts(db, out['payment_id'])
    assert (r.provider, r.status, r.ebarimt_id) == ('ONTIME', 'SENT', body['ebarimt_id'])
    assert jobs(db, out['payment_id']) == []


# ── 3. Асаалттай + веб касс (толгойгүй) бэлэн → msgbill хэвээр ─────────────
def test_setting_on_web_cash_without_header_still_msgbill(db):
    s, u = parking(db), user(db, 'online.operator')
    enable(db, s.site_id)
    out = cash(db, s, u)
    p = db.get(M.Payment, out['payment_id'])
    assert p.source is None
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING') and len(jobs(db, p.id)) == 1
    audit = db.query(M.AuditLog).filter_by(action='CASH_PAYMENT').one()
    assert audit.detail['source'] is None and audit.detail['app_version'] is None


# ── 4. Асаалттай + POS аппын бэлэн (толгойтой) → EXTERNAL ──────────────────
def test_setting_on_pos_app_cash_with_header_is_external(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    out = cash(db, s, u, app_version='1.4.0')
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and p.raw_payload['app_version'] == '1.4.0'
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('POS-TERMINAL', 'EXTERNAL') and jobs(db, p.id) == []
    assert db.query(M.AuditLog).filter_by(action='CASH_PAYMENT').one().detail['source'] == 'POSAPP'


def test_site_off_override_does_not_affect_other_sites(db):
    """Терминал нь хэвлэдэггүй НЭГ зогсоолд унтраавал зөвхөн тэр зогсоолд систем баримт
    гаргана; бусад (ерөнхий ON) хэвээр EXTERNAL."""
    s, u = parking(db), user(db)
    other = parking(db, code='MONNIS', plate='5678УБА')
    enable(db, s.site_id, on=False)           # зөвхөн Хангарьдад унтраав
    out = cash(db, s, u, app_version='1.4.0')
    [r] = receipts(db, out['payment_id'])
    assert r.provider == 'MSGBILL' and len(jobs(db, out['payment_id'])) == 1
    out2 = cash(db, other, u, app_version='1.4.0')
    assert receipts(db, out2['payment_id'])[0].status == 'EXTERNAL'


# ── 5. Аргын шүүлтүүр ──────────────────────────────────────────────────────
def test_method_filter_card_only(db):
    s1, u = parking(db), user(db)
    s2 = M.ParkingSession(site_id=s1.site_id, plate_number='9999УБА',
                          entry_time=datetime.utcnow() - timedelta(minutes=59), status='AWAITING_PAYMENT')
    db.add(s2); db.flush(); db.refresh(s2)
    payment_wait.begin_wait(db, s2, SL.session_fee_info(db, s2), datetime.utcnow()); db.commit()
    enable(db, s1.site_id, methods='CARD')
    c_out = card(db, s1, u)
    k_out = cash(db, s2, u, app_version='1.4.0')
    assert receipts(db, c_out['payment_id'])[0].status == 'EXTERNAL'
    assert receipts(db, k_out['payment_id'])[0].provider == 'MSGBILL'


def test_methods_are_validated_and_normalized(db):
    vals = PRULES.clamp_values(A.EBARIMT_KEY, {'pos_terminal_methods': ' cash, card '})
    assert vals == {'pos_terminal_methods': 'CARD,CASH'}
    site = 'site-x'
    A.set_site_rules(db, A.EBARIMT_KEY, site, {'pos_terminal_issues': True,
                                                'pos_terminal_methods': 'QR'}, 'admin')
    db.commit()
    # Танигдахгүй арга хадгалагдахгүй → асаалттай үед анхдагч CARD+CASH
    assert A.pos_terminal_receipt_methods(db, site) == {'CARD', 'CASH'}
    A.set_site_rules(db, A.EBARIMT_KEY, site, {'pos_terminal_methods': 'CASH'}, 'admin'); db.commit()
    assert A.pos_terminal_receipt_methods(db, site) == {'CASH'}


def test_settings_endpoint_is_per_site_and_audited(db):
    from app.routers import admin_router as AR
    s, u = parking(db), user(db, 'admin')
    out = AR.payment_rules_site_save(s.site_id, {'ebarimt_rules': {
        'pos_terminal_issues': True, 'pos_terminal_methods': 'cash,card'}}, db, u)
    row = {r['key']: r for r in out['rules'] if r['group'] == A.EBARIMT_KEY}
    assert row['pos_terminal_issues']['value'] is True and row['pos_terminal_issues']['source'] == 'site'
    assert row['pos_terminal_methods']['value'] == 'CARD,CASH'
    assert row['pos_terminal_methods']['choices']
    log = db.query(M.AuditLog).filter_by(entity='payment_rules', entity_id=s.site_id).one()
    assert log.detail['ebarimt_rules']['pos_terminal_issues'] is True
    # Зогсоолд УНТРААХ — ерөнхий (анхдагч ON) хөндөгдөхгүй
    out = AR.payment_rules_site_save(s.site_id, {'ebarimt_rules': {'pos_terminal_issues': False}}, db, u)
    row = {r['key']: r for r in out['rules'] if r['group'] == A.EBARIMT_KEY}
    assert row['pos_terminal_issues']['value'] is False and row['pos_terminal_issues']['global_value'] is True
    assert A.get_rules(db, A.EBARIMT_KEY)['pos_terminal_issues'] is True
    assert A.pos_terminal_receipt_methods(db, s.site_id) == set()


# ── 6. Retry/бөөн нөхөлт/ажил EXTERNAL-ийг хөндөхгүй ───────────────────────
def external_payment(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    out = card(db, s, u)
    return db.get(M.Payment, out['payment_id']), u, s


def test_retry_refuses_external_with_409(db):
    p, u, _ = external_payment(db)
    with pytest.raises(HTTPException) as e:
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))
    assert e.value.status_code == 409 and 'POS терминал' in e.value.detail
    db.rollback()
    # Дүрэм үйлчилж байхад «Систем баримт гаргах» нь «терминал хэвлээгүй» баталгаажуулалтгүй
    # бол татгалзана (хуучин клиент/скрипт санамсаргүй гаргахгүй)
    with pytest.raises(HTTPException) as e:
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True}, db, u))
    assert e.value.status_code == 409 and 'no_terminal_receipt' in e.value.detail
    db.rollback()
    assert receipts(db, p.id)[0].status == 'EXTERNAL' and db.query(M.FinancialJob).count() == 0
    assert J.retry_payment_receipts(db, p)['external'] is True
    db.rollback()


def test_reissue_with_no_terminal_confirmation_while_rule_on(db):
    """2026-10-14: дүрэм ерөнхийдөө асаалттай — нэг борлуулалтын төлөө зогсоолын тохиргоог
    унтраахгүй. «Терминал хэвлээгүй» баталгаажуулалттай бол систем баримт гаргана."""
    p, u, s = external_payment(db)
    res = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True,
                                                       'no_terminal_receipt': True}, db, u))
    assert res['pending'] is True
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING')
    [j] = jobs(db, p.id)
    assert j.payload['idempotency_key'] == f'external-{r.id}'
    assert A.pos_terminal_receipt_methods(db, s.site_id) == {'CARD', 'CASH'}     # тохиргоо хөндөгдөөгүй
    log = db.query(M.AuditLog).filter_by(action='EBARIMT_RETRY', entity_id=p.id).one()
    assert log.detail['no_terminal_receipt'] is True
    assert db.query(M.AuditLog).filter_by(action='EBARIMT_EXTERNAL_REISSUE').one().detail['no_terminal_receipt']


def test_reissue_only_after_setting_off_and_explicit_request(db):
    p, u, s = external_payment(db)
    enable(db, s.site_id, on=False)
    with pytest.raises(HTTPException) as e:        # тодорхой хүсэлтгүй бол унтраасан ч 409
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))
    assert e.value.status_code == 409
    db.rollback()
    res = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True}, db, u))
    assert res['pending'] is True
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING')
    [j] = jobs(db, p.id)
    assert j.receipt_id == r.id and j.payload['provider'] == 'MSGBILL' and j.status == 'PENDING'
    # Аудитад товч дарсан АДМИН (өмнө нь 'system' гэж бичигддэг байв — хяналт тэмдэглэл b)
    [log] = db.query(M.AuditLog).filter_by(action='EBARIMT_EXTERNAL_REISSUE').all()
    assert log.username == u.username


def test_worker_never_sends_external_receipt(db, monkeypatch):
    p, _, _ = external_payment(db)
    [r] = receipts(db, p.id)
    # Хэн нэгэн гараар ажил үүсгэсэн ч (DB засвар г.м) сувагт илгээхгүй
    db.add(M.FinancialJob(job_key=f'receipt:{r.id}', kind='RECEIPT', payment_id=p.id, receipt_id=r.id,
                          payload={'provider': 'MSGBILL', 'amount': '1000', 'vat': '91',
                                   'method': 'CARD', 'customer_tin': None, 'receiver_type': 'CITIZEN',
                                   'description': 'x', 'idempotency_key': 'k'}))
    db.commit()
    create = AsyncMock(side_effect=AssertionError('msgbill POST for EXTERNAL'))
    monkeypatch.setattr(J.msgbill, 'create_receipt', create)
    asyncio.run(J.run_once(factory=lambda: Session(db.bind, autoflush=False)))
    db.expire_all()
    assert create.await_count == 0
    assert db.query(M.FinancialJob).one().status == 'DONE'
    assert receipts(db, p.id)[0].status == 'EXTERNAL'


def test_reissue_reuses_job_left_done_by_worker_guard(db, monkeypatch):
    p, u, s = external_payment(db)
    [r] = receipts(db, p.id)
    db.add(M.FinancialJob(job_key=f'receipt:{r.id}', kind='RECEIPT', payment_id=p.id, receipt_id=r.id,
                          payload={'provider': 'MSGBILL'}, status='DONE'))
    db.commit()
    enable(db, s.site_id, on=False)
    asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True}, db, u))
    [j] = jobs(db, p.id)                            # job_key unique — шинэ мөр биш, дахин ашиглав
    assert j.status == 'PENDING' and j.payload['idempotency_key'] == f'external-{r.id}'
    assert '1234УБА' in j.payload['description']


def test_bulk_retry_never_selects_or_reissues_external(db, monkeypatch):
    """Бөөн нөхөлт (vat-retry-failed) зөвхөн FAILED/REVIEW мөрийг сонгодог тул ЦЭВЭР
    EXTERNAL төлбөр нэр дэвшигч ч болохгүй; FAILED+EXTERNAL холимог (тохиргоо солигдсон
    үеийн хуучин мөр г.м) төлбөр сонгогдсон ч msgbill-ээр гаргалгүй алгасна."""
    import app.database as D
    p, u, s = external_payment(db)
    dry = asyncio.run(R.vat_retry_failed({'dry': True}, db, u))
    assert dry['candidates'] == 0                      # цэвэр EXTERNAL — сонгогдохгүй
    db.add(M.VatReceipt(payment_id=p.id, session_id=s.id, amount=Decimal('100'),
                        vat_amount=Decimal('9'), provider='MSGBILL', status='FAILED',
                        receipt_url='msgbill 500'))
    db.commit()
    dry = asyncio.run(R.vat_retry_failed({'dry': True}, db, u))
    assert dry['candidates'] == 1
    monkeypatch.setattr(D, 'SessionLocal', lambda: Session(db.bind, autoflush=False))
    R._bulk_job.clear()
    R._bulk_job.update({'running': True, 'total': 1, 'done': 0, 'ok': 0, 'skipped': 0,
                        'failed': 0, 'errors': {}, 'username': 't'})
    monkeypatch.setattr(asyncio, 'sleep', AsyncMock())
    asyncio.run(R._run_bulk_retry([p.id]))
    assert R._bulk_job['skipped'] == 1 and R._bulk_job['failed'] == 0
    assert R._bulk_job['external_skipped'] == 1
    db.expire_all()
    assert db.query(M.FinancialJob).count() == 0
    assert sorted(r.status for r in receipts(db, p.id)) == ['EXTERNAL', 'FAILED']
    R._bulk_job.clear()


def test_cancel_skips_external(db):
    p, _, _ = external_payment(db)
    res = asyncio.run(PR.cancel_ebarimt(db, p, 'давхар'))
    assert res['ok'] is False and res['external'] is True and 'Ontime' in res['error']
    assert receipts(db, p.id)[0].status == 'EXTERNAL'


def test_attach_adopts_external_row(db):
    p, u, _ = external_payment(db)
    res = asyncio.run(PR.attach_external_ebarimt(p.id, {'ebarimt_id': 'ONTIME-DDTD-1',
                                                        'ebarimt_provider': 'ONTIME'}, db, u))
    assert res['ok'] and res['ebarimt_id'] == 'ONTIME-DDTD-1'
    [r] = receipts(db, p.id)                       # шинэ мөр нэмээгүй — EXTERNAL нь SENT болов
    assert (r.status, r.provider, r.ebarimt_id) == ('SENT', 'ONTIME', 'ONTIME-DDTD-1')


# ── 7. primary_receipt / хэвлэх ────────────────────────────────────────────
def _rec(status, mins, ebarimt_id=None):
    return NS(status=status, ebarimt_id=ebarimt_id, created_at=datetime(2026, 10, 13) + timedelta(minutes=mins))


def test_primary_receipt_ranks_external():
    ext, sent = _rec('EXTERNAL', 0), _rec('SENT', 5, 'D1')
    cancelled, pending, failed = _rec('CANCELLED', -5, 'OLD'), _rec('PENDING', -1), _rec('FAILED', -2)
    assert choose_primary([ext, sent]) is sent              # жинхэнэ ДДТД-тэй мөр давуу
    assert choose_primary([cancelled, ext]) is ext          # цуцалсан давхардлаас өмнө
    assert choose_primary([pending, failed, ext]) is ext    # эцсийн төлөв — түр зуурынхаас өмнө


def test_print_payload_and_pos_receipts_for_external(db):
    p, u, _ = external_payment(db)
    rows = PR.pos_recent_receipts(None, 10, db, u)
    assert rows[0]['receipt']['ebarimt_status'] == 'EXTERNAL'
    assert 'ДДТД: -' not in rows[0]['print_data']['lines']
    single = PR.payment_print(p.id, db, u)
    assert single['receipt']['ebarimt_issuer'] == 'POS_TERMINAL'


def test_print_returns_ddtd_once_outbox_delivers(db):
    """Зогсоолд унтраалттай: confirm-ийн мөчид PENDING, outbox дууссаны дараа /print
    болон /pos/receipts ДДТД-тэй (блоклох хүлээлтгүй — апп дахин татна)."""
    s, u = parking(db), user(db)
    enable(db, s.site_id, on=False)
    out = card(db, s, u)
    assert out['receipt']['ebarimt_status'] == 'PENDING' and out['ebarimt_id'] is None
    [r] = receipts(db, out['payment_id'])
    r.ebarimt_id, r.status, r.lottery_code = '0301010000005595', 'SENT', 'HV 1'
    db.commit()
    again = PR.payment_print(out['payment_id'], db, u)
    assert again['ebarimt_id'] == '0301010000005595'
    assert 'ДДТД: 0301010000005595' in again['print_data']['lines']
    assert PR.pos_recent_receipts(None, 10, db, u)[0]['receipt']['ebarimt_status'] == 'SENT'


# ── 8. Өрийн төлөлт (Нөхөн төлбөр) ─────────────────────────────────────────
def debt(db, site_id):
    c = M.Compensation(site_id=site_id, plate_number='1234УБА', amount=Decimal('5000'))
    db.add(c); db.commit()
    return c


def test_compensation_pay_from_pos_app_is_external(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    c = debt(db, s.site_id)
    out = asyncio.run(C.pay_compensation(c.id, {'method': 'CASH'}, db, u, req('1.4.0')))
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and out['receipt_status'] == 'EXTERNAL'
    assert jobs(db, p.id) == []


def test_compensation_pay_from_web_keeps_msgbill(db):
    s, u = parking(db), user(db, 'operator')
    enable(db, s.site_id)
    c = debt(db, s.site_id)
    out = asyncio.run(C.pay_compensation(c.id, {'method': 'CASH'}, db, u, req()))
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POS' and out['receipt_status'] == 'PENDING'
    assert receipts(db, p.id)[0].provider == 'MSGBILL' and len(jobs(db, p.id)) == 1


def test_combined_payment_with_debts_is_one_terminal_receipt(db):
    """Терминал НИЙТ дүнгээр НЭГ баримт хэвлэдэг — өр тус бүрийн мөр хуваахгүй."""
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    c = debt(db, s.site_id)
    p = M.Payment(session_id=s.id, provider='POS', payment_method='CARD', source='POSAPP',
                  sender_invoice_no='combo-1', amount=Decimal('6000'), vat_amount=Decimal('545'),
                  status='PENDING', fee_snapshot={'parking_amount': 1000.0,
                                                  'debts': [{'id': c.id, 'amount': 5000.0}]})
    db.add(p); db.commit()
    asyncio.run(PR._finalize_paid(db, p))
    [r] = receipts(db, p.id)
    assert r.status == 'EXTERNAL' and r.amount == Decimal('6000')
    assert db.get(M.Compensation, c.id).status == 'PAID'
    assert db.query(M.FinancialJob).count() == 0


# ── 9. Тайлан: ТЕГ тулгалт, Excel ──────────────────────────────────────────
def test_reconcile_explains_terminal_rows():
    t = datetime(2026, 8, 24, 10, 0, 0)
    left = [{'dt': t, 'amount': 2000.0, 'src': 'f', 'ddtd': 'X10002990'}]
    probe = [{'id': 'p1', 'paid_at': t, 'amount': 2000.0, 'site_id': 's', 'plate': '1',
              'site_name': 'Хангарьд', 'provider': 'POS', 'method': 'CARD', 'receipts': [],
              'terminal': True}]
    [row] = V.explain_unmatched_tax(left, probe, 0, 3, set(), None)
    assert row['verdict'] == 'TERMINAL' and 'TERMINAL' in V.VERDICTS
    probe[0]['terminal'] = False
    assert V.explain_unmatched_tax(left, probe, 0, 3, set(), None)[0]['verdict'] == 'PAYMENT_NO_RECEIPT'


def test_excel_counts_external_separately():
    from openpyxl import load_workbook
    mk = lambda st, prov, amt, url=None: NS(status=st, provider=prov, amount=amt, vat_amount=0,
                                            ebarimt_id=None if st == 'EXTERNAL' else 'D', lottery_code=None,
                                            created_at=datetime(2026, 10, 13), customer_tin=None,
                                            receipt_url=url)
    rows = [(mk('SENT', 'MSGBILL', 1000), '1', 'A'),
            (mk('EXTERNAL', 'POS-TERMINAL', 2000, J.TERMINAL_NOTE), '2', 'A')]
    body = asyncio.run(_collect(X.vat_receipts_excel(rows)))   # StreamingResponse
    ws = load_workbook(BytesIO(body)).active
    values = [[c.value for c in r] for r in ws.iter_rows()]
    assert values[2][8] == 'POS терминал гаргасан' and values[2][7] == 'POS терминал (Ontime)'
    assert values[2][10] == J.TERMINAL_NOTE
    assert 'POS терминал гаргасан: 1 мөр' in values[-1][10]


async def _collect(resp):
    out = b''
    async for chunk in resp.body_iterator:
        out += chunk if isinstance(chunk, bytes) else chunk.encode()
    return out


# ══ 2026-10-13 хяналтын засварууд ══════════════════════════════════════════
def another_session(db, site_id, plate):
    s = M.ParkingSession(site_id=site_id, plate_number=plate,
                         entry_time=datetime.utcnow() - timedelta(minutes=59), status='AWAITING_PAYMENT')
    db.add(s); db.flush(); db.refresh(s)
    payment_wait.begin_wait(db, s, SL.session_fee_info(db, s), datetime.utcnow()); db.commit()
    return s


# ── #1: толгойгүй бэлэн — зогсоолын «POS аппын хэрэглэгчид» жагсаалт ──────────
def test_user_list_marks_headerless_cash_as_pos_app(db):
    """Апп X-App-Version илгээдэг нь прод-д батлагдаагүй — давхардлын 62% нь бэлэн.
    Жагсаалтын хэрэглэгчийн толгойгүй бэлэн → EXTERNAL; веб касс (sogii) → msgbill."""
    s1, pos_user, web = parking(db), user(db, 'pos.khangarid'), user(db, 'sogii')
    A.set_site_rules(db, A.EBARIMT_KEY, s1.site_id,
                     PRULES.clamp_values(A.EBARIMT_KEY, {'pos_terminal_issues': True,
                                                         'pos_terminal_users': ' POS.khangarid ; other'}),
                     'admin')
    db.commit()
    out = cash(db, s1, pos_user)                                   # толгойгүй
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and p.raw_payload['pos_app_via'] == 'user_list'
    [r] = receipts(db, p.id)
    assert r.status == 'EXTERNAL' and jobs(db, p.id) == []
    log = db.query(M.AuditLog).filter_by(action='CASH_PAYMENT', entity_id=p.id).one()
    assert log.detail['pos_app_via'] == 'user_list' and log.detail['app_version'] is None
    s2 = another_session(db, s1.site_id, '5555УБА')
    out2 = cash(db, s2, web)                                       # веб касс — хуучин зан
    p2 = db.get(M.Payment, out2['payment_id'])
    assert p2.source is None and receipts(db, p2.id)[0].provider == 'MSGBILL'
    assert len(jobs(db, p2.id)) == 1
    # Өрийн төлөлт (Нөхөн төлбөр) — мөн жагсаалтаар
    c = debt(db, s1.site_id)
    out3 = asyncio.run(C.pay_compensation(c.id, {'method': 'CASH'}, db, pos_user, req()))
    assert out3['receipt_status'] == 'EXTERNAL'
    assert db.get(M.Payment, out3['payment_id']).source == 'POSAPP'


def test_user_list_is_global(db):
    """2026-10-14: POS аппын хэрэглэгчийг зогсоолоос ҮЛ хамааран таньна — A зогсоолын
    жагсаалтад бичсэн хэрэглэгч B зогсоолд толгойгүй бэлэн авсан ч POS апп."""
    s1, u = parking(db), user(db, 'pos.khangarid')
    s2 = parking(db, code='OTHER', plate='7777УБА')
    A.set_site_rules(db, A.EBARIMT_KEY, s1.site_id, {'pos_terminal_users': 'pos.khangarid'}, 'admin')
    db.commit()
    out = cash(db, s2, u)
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and p.raw_payload['pos_app_via'] == 'user_list'
    assert receipts(db, p.id)[0].status == 'EXTERNAL'


# ── #5 (2026-10-14): «терминал баримт гаргадаг» ЕРӨНХИЙ анхдагч + зогсоолын давхарга ──
def test_terminal_setting_is_global_with_site_override(db):
    from app.routers import admin_router as AR
    s, admin = parking(db), user(db, 'admin')
    other = parking(db, code='OTHER', plate='7777УБА')
    rows = {r['key']: r for r in PRULES.global_report(db)['rules'] if r['group'] == A.EBARIMT_KEY}
    assert rows['pos_terminal_issues']['value'] is True and not rows['pos_terminal_issues'].get('site_only')
    assert rows['pos_terminal_issues']['per_site'] is True
    # Ерөнхийг УНТРААВАЛ тохиргоогүй зогсоолд үйлчилнэ; зогсоолд тодорхой ON хэвээр
    A.set_site_rules(db, A.EBARIMT_KEY, s.site_id, {'pos_terminal_issues': True}, 'admin'); db.commit()
    out = AR.payment_rules_site_save('global', {'ebarimt_rules': {
        'pos_terminal_issues': False, 'pos_terminal_users': 'pos.khangarid',
        'pos_terminal_methods': 'card'}}, db, admin)
    rows = {r['key']: r for r in out['rules'] if r['group'] == A.EBARIMT_KEY}
    assert rows['pos_terminal_issues']['value'] is False
    assert rows['pos_terminal_users']['value'] == 'pos.khangarid'
    assert rows['pos_terminal_methods']['value'] == 'CARD'
    assert A.pos_terminal_receipt_methods(db, other.site_id) == set()          # ерөнхийг өвлөв
    assert A.pos_terminal_receipt_methods(db, s.site_id) == {'CARD'}           # зогсоолд ON, арга ерөнхийгөөс
    assert A.pos_app_users(db) == {'pos.khangarid'}
    # Ерөнхийг буцааж асаавал зогсоолын OFF давхарга дарна
    AR.payment_rules_site_save('global', {'ebarimt_rules': {'pos_terminal_issues': True,
                                                             'pos_terminal_methods': 'CARD,CASH'}}, db, admin)
    AR.payment_rules_site_save(other.site_id, {'ebarimt_rules': {'pos_terminal_issues': False}}, db, admin)
    assert A.pos_terminal_receipt_methods(db, other.site_id) == set()
    assert A.pos_terminal_receipt_methods(db, s.site_id) == {'CARD', 'CASH'}
    # «Ерөнхий рүү буцаах» (null) — давхаргаас хасна
    AR.payment_rules_site_save(other.site_id, {'ebarimt_rules': {'pos_terminal_issues': None}}, db, admin)
    assert A.pos_terminal_receipt_methods(db, other.site_id) == {'CARD', 'CASH'}


def test_existing_per_site_on_rows_stay_valid(db):
    """Продын 7 зогсоолын ON мөр (users-тай) — шинэ ерөнхий ON-той зөрчилгүй, хүчинтэй."""
    s, u = parking(db), user(db, 'pos.khangarid')
    A.set_site_rules(db, A.EBARIMT_KEY, s.site_id, {'pos_terminal_issues': True,
                                                    'pos_terminal_methods': 'CARD,CASH',
                                                    'pos_terminal_users': 'pos.khangarid'}, 'admin')
    db.commit()
    assert A.pos_terminal_receipt_methods(db, s.site_id) == {'CARD', 'CASH'}
    out = cash(db, s, u)                                            # толгойгүй — жагсаалтаар
    assert receipts(db, out['payment_id'])[0].status == 'EXTERNAL'


def test_conflict_warns_when_site_with_pos_traffic_is_off(db):
    """Зөрчлийн шалгалт: POS аппын төлбөртэй зогсоолд дүрэм унтраалттай бол анхааруулна."""
    s, u = parking(db), user(db)
    site = db.get(M.ParkingSite, s.site_id)
    titles = lambda: [c['title'] for c in PRULES.check_conflicts(db, site)]   # noqa: E731
    enable(db, s.site_id, on=False)
    assert not any('POS аппын төлбөртэй' in t for t in titles())          # POS аппын төлбөр алга
    card(db, s, u)
    assert any('POS аппын төлбөртэй' in t for t in titles())
    enable(db, s.site_id, methods='CASH')                                   # карт хасагдсан
    assert any('POS аппын төлбөртэй' in t for t in titles())
    enable(db, s.site_id)
    assert not any('POS аппын төлбөртэй' in t for t in titles())


# ── #2: давхардал гэж цуцалсны дараа «Дахин үүсгэх» давхардлыг сэргээхгүй ───
def _ddtd(seq, register):
    return f'0152000200900009726{seq:06d}{register}'      # 33 орон, сүүлийн 8 = касс


MSGBILL_POS, ONTIME_POS = '10045952', '10002990'


def historical_msgbill(db, s, u, *, method='CASH', seq=1, cancelled=False, mark=False):
    """Тохиргоо нэвтрэхээс ӨМНӨХ POS аппын төлбөр (source NULL) + msgbill баримт."""
    p = M.Payment(session_id=s.id, provider='CASH' if method == 'CASH' else 'POS',
                  payment_method=method, sender_invoice_no=f'hist-{seq}', amount=Decimal('2000'),
                  vat_amount=Decimal('182'), status='PAID', paid_at=datetime.utcnow(),
                  cashier_id=u.id, raw_payload={})
    db.add(p); db.flush()
    r = M.VatReceipt(payment_id=p.id, session_id=s.id, amount=p.amount, vat_amount=p.vat_amount,
                     provider='MSGBILL', provider_ref=f'rcp_{seq}', ebarimt_id=_ddtd(seq, MSGBILL_POS),
                     status='CANCELLED' if cancelled else 'SENT')
    if mark:
        r.ddtd_note = 'DUPLICATE_OF_POS_TERMINAL: POS терминалын баримттай давхардсан тул цуцлав'
    db.add(r); db.commit()
    return p, r


def cancel_dups(db, u, ids, teg=None):
    """API-ийн арын job ба CLI-ийн НЭГ зам (vat_duplicates.cancel_one) — нэгтгэлийн дараа
    давхардлын цуцлалт зөвхөн үүгээр явна."""
    return [asyncio.run(VD.cancel_one(db, pid, note='ТЕГ тулгалт: Ontime давхардал',
                                      username=u.username,
                                      teg={'teg_ddtd': (teg or {}).get(pid)}))
            for pid in ids]


def test_duplicate_cancel_records_terminal_receipt_and_retry_is_noop(db, monkeypatch):
    s, u = parking(db), user(db)
    p, r = historical_msgbill(db, s, u)
    cancel = AsyncMock(return_value={'state': 'CANCELLED'})
    monkeypatch.setattr(PR.msgbill, 'cancel_receipt', cancel)
    teg = _ddtd(77, ONTIME_POS)
    res = cancel_dups(db, u, [p.id], {p.id: teg})
    assert res[0]['ok'] is True and cancel.await_count == 1
    tr = res[0]['terminal_receipt']
    assert (tr['status'], tr['ebarimt_id'], tr['provider']) == ('SENT', teg, 'ONTIME')
    db.expire_all()
    rows = {x.provider: x for x in receipts(db, p.id)}
    assert rows['MSGBILL'].status == 'CANCELLED'
    assert rows['MSGBILL'].raw['cancel']['reason'] == 'DUPLICATE_OF_POS_TERMINAL'
    assert rows['MSGBILL'].raw['cancel']['teg_ddtd'] == teg
    assert (rows['ONTIME'].status, rows['ONTIME'].ebarimt_id) == ('SENT', teg)
    assert rows['ONTIME'].amount == p.amount
    # «Дахин үүсгэх» — терминалын ДДТД-тэй тул юу ч үүсгэхгүй (msgbill-гүй)
    out = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))
    assert out['ok'] is True and out['ebarimt_id'] == teg
    assert len(receipts(db, p.id)) == 2 and db.query(M.FinancialJob).count() == 0
    # Хэвлэлт терминалын ДДТД-г харуулна (цуцлагдсан msgbill-ийн биш)
    assert f'ДДТД: {teg}' in PR.payment_print(p.id, db, u)['print_data']['lines']
    # Дахин ажиллуулбал (ижил ТЕГ файл) — терминалын мөрийг ХӨНДӨХГҮЙ
    again = cancel_dups(db, u, [p.id], {p.id: teg})
    assert again[0]['skipped'] == 'already_cancelled' and 'terminal_receipt' not in again[0]
    assert cancel.await_count == 1
    db.expire_all()
    assert {x.provider: x.status for x in receipts(db, p.id)} == {'MSGBILL': 'CANCELLED', 'ONTIME': 'SENT'}


def test_duplicate_cancel_with_untrusted_ddtd_records_external(db, monkeypatch):
    """ТЕГ-ийн «давхар» мөр манай msgbill-ийн КАССтай (хос солигдсон) — терминалын ДДТД
    гэж итгэхгүй, ДДТД-гүй EXTERNAL бүртгэнэ; дахин гаргахыг зөвхөн ил зөвшөөрлөөр."""
    s, u = parking(db), user(db)
    p, r = historical_msgbill(db, s, u, method='CARD')
    monkeypatch.setattr(PR.msgbill, 'cancel_receipt', AsyncMock(return_value={'state': 'CANCELLED'}))
    res = cancel_dups(db, u, [p.id], {p.id: _ddtd(78, MSGBILL_POS)})
    assert res[0]['terminal_receipt']['status'] == 'EXTERNAL'
    db.expire_all()
    ext = [x for x in receipts(db, p.id) if x.status == 'EXTERNAL']
    assert len(ext) == 1 and ext[0].ebarimt_id is None and ext[0].raw['teg']['untrusted'] is True
    with pytest.raises(HTTPException) as e:                        # энгийн retry
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))
    assert e.value.status_code == 409
    db.rollback()
    with pytest.raises(HTTPException) as e:                        # тохиргоо унтраалттай ч давхардал
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True}, db, u))
    assert e.value.status_code == 409 and 'ДАВХАРДСАН' in e.value.detail
    db.rollback()
    assert db.query(M.AuditLog).filter_by(action='EBARIMT_RETRY_REFUSED').count() == 1
    # Давхардлыг зөвшөөрсөн ч POS терминалын дүрэм (ерөнхий ON) үйлчилж байгаа тул
    # «терминал хэвлээгүй» баталгаажуулалтгүйгээр гаргахгүй
    with pytest.raises(HTTPException) as e:
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True,
                                                     'allow_duplicate': True}, db, u))
    assert e.value.status_code == 409 and 'no_terminal_receipt' in e.value.detail
    db.rollback()
    out = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'reissue_external': True, 'allow_duplicate': True,
                                                       'no_terminal_receipt': True}, db, u))
    assert out['pending'] is True
    assert {x.provider for x in receipts(db, p.id) if x.status == 'PENDING'} == {'MSGBILL'}


def test_retry_refuses_duplicate_cancelled_without_terminal_row(db):
    """Цуцлалт өөр замаар (CLI/job) хийгдэж терминалын мөр бүртгэгдээгүй ч тэмдэг хангалттай —
    хяналтын давталт: CANCELLED msgbill + retry {} → өмнө нь шинэ msgbill PENDING + ажил."""
    s, u = parking(db), user(db)
    p, _ = historical_msgbill(db, s, u, cancelled=True, mark=True)
    with pytest.raises(HTTPException) as e:
        asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))
    assert e.value.status_code == 409 and 'ДАВХАРДСАН' in e.value.detail
    db.rollback()
    assert len(receipts(db, p.id)) == 1 and db.query(M.FinancialJob).count() == 0
    out = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {'allow_duplicate': True}, db, u))
    assert out['pending'] is True and len(jobs(db, p.id)) == 1


# ── #4: бүгд цуцлагдсан + тохиргоо асаалттай → EXTERNAL бүртгээд 200 ──────────
def test_retry_all_cancelled_with_setting_on_records_external_and_returns_200(db):
    s, u = parking(db), user(db)
    p, _ = historical_msgbill(db, s, u, method='CARD', cancelled=True)   # provider POS
    enable(db, s.site_id)
    out = asyncio.run(PR.retry_ebarimt_endpoint(p.id, {}, db, u))      # 409 биш
    assert out['created'] is True and out['external'] is True
    assert sorted(x.status for x in receipts(db, p.id)) == ['CANCELLED', 'EXTERNAL']
    assert db.query(M.FinancialJob).count() == 0
    assert db.query(M.AuditLog).filter_by(action='EBARIMT_RETRY', entity_id=p.id).count() == 1


# ── тэмдэглэл d: өрийн төлөлтийн дахин гаргалтын тайлбар ─────────────────────
def test_reissue_of_debt_payment_keeps_debt_description(db):
    s, u = parking(db), user(db)
    enable(db, s.site_id)
    c = debt(db, s.site_id)
    out = asyncio.run(C.pay_compensation(c.id, {'method': 'CASH'}, db, u, req('1.4.0')))
    enable(db, s.site_id, on=False)
    asyncio.run(PR.retry_ebarimt_endpoint(out['payment_id'], {'reissue_external': True}, db, u))
    [j] = jobs(db, out['payment_id'])
    assert j.payload['description'] == 'Зогсоолын өр · 1234УБА'


# ── #3: ТЕГ тулгалт — терминалын баримт ОЛДСОН/ДУТУУ ─────────────────────────
def _paid(db, s, u, at, amount, *, external, seq):
    p = M.Payment(session_id=s.id, provider='POS', payment_method='CARD', source='POSAPP',
                  sender_invoice_no=f'rc-{seq}', amount=Decimal(amount), vat_amount=Decimal('0'),
                  status='PAID', paid_at=at, cashier_id=u.id, raw_payload={})
    db.add(p); db.flush()
    r = M.VatReceipt(payment_id=p.id, session_id=s.id, amount=p.amount, vat_amount=p.vat_amount)
    if external:
        r.provider, r.status, r.receipt_url = 'POS-TERMINAL', 'EXTERNAL', J.TERMINAL_NOTE
    else:
        r.provider, r.status, r.ebarimt_id = 'MSGBILL', 'SENT', _ddtd(seq, MSGBILL_POS)
    db.add(r)
    return p


def test_vat_reconcile_confirms_and_lists_missing_terminal_receipts(db, monkeypatch):
    from starlette.datastructures import UploadFile
    from sqlalchemy.orm import sessionmaker
    import app.database as D
    # API тулгалтыг thread дээр ӨӨРИЙН session-ээр хийдэг (vat_duplicates.reconcile_db)
    monkeypatch.setattr(D, 'SessionLocal', sessionmaker(bind=db.get_bind(), autoflush=False))
    s, u = parking(db), user(db)
    base = datetime(2026, 10, 12, 4, 0, 0)
    # 5-р (сүүлийн) манай баримт — файлын сүүлийн мөрөөс tol-оор өмнөх хүртэл л терминалын
    # нотолгооны хүрээ (хяналт 2: сүүлчийн борлуулалтын хожуу мөр файлд орохгүй байж болно)
    ours = [_paid(db, s, u, base + timedelta(minutes=10 * i), 1000 + 100 * i, external=False, seq=i)
            for i in range(5)]
    term = [_paid(db, s, u, base + timedelta(minutes=10 * i + 5), 3000 + 100 * i, external=True,
                  seq=10 + i) for i in range(4)]
    missing = _paid(db, s, u, base + timedelta(minutes=17), 4500, external=True, seq=20)
    db.commit()
    loc = timedelta(hours=8)          # ТЕГ: msgbill ба терминалын касс хоёулаа УБ локал цагаар
    lines = ['ТТД,ДДТД,Огноо,Нийт дүн']
    for i, p in enumerate(ours):
        lines.append(f'7524322,{_ddtd(100 + i, MSGBILL_POS)},{(p.paid_at + loc):%Y-%m-%d %H:%M:%S},{p.amount}')
    for i, p in enumerate(term):
        lines.append(f'7524322,{_ddtd(200 + i, ONTIME_POS)},{(p.paid_at + loc):%Y-%m-%d %H:%M:%S},{p.amount}')
    data = '\n'.join(lines).encode()

    def run(excel=False):
        return asyncio.run(R.vat_reconcile(
            file=UploadFile(file=BytesIO(data), filename='teg.csv'), tz_shift=None, tol=3,
            excel=excel, tenant_id=None, col_ddtd=None, col_dt=None, col_amount=None, db=db, user=u))

    res = run()
    assert res['ours_total'] == 5 and res['matched'] == 5           # EXTERNAL «манайх»-т ороогүй
    assert res['unmatched_ours_total'] == 0
    assert res['terminal_issued'] == 5 and res['terminal_confirmed'] == 4
    assert res['terminal_missing_total'] == 1
    assert res['terminal_missing'][0]['payment_id'] == missing.id
    assert res['terminal_missing'][0]['status'] == 'EXTERNAL'
    assert res['unmatched_tax_total'] == 0                          # Ontime мөр бүр тайлагдсан
    assert res['tax_verdicts'] == {'TERMINAL': 4}
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(asyncio.run(_collect(run(excel=True)))))
    rows = [[c.value for c in r] for r in wb['Манайд бий - ТЕГ-д алга'].iter_rows()]
    assert len(rows) == 2 and rows[1][5] == 'EXTERNAL' and 'ТЕГ-т баримт АЛГА' in rows[1][7]


# ── #6: Бичилт — баримтын төлөв/гаргагч ─────────────────────────────────────
def test_transactions_report_shows_terminal_receipt_state(db):
    from openpyxl import load_workbook
    p, u, s = external_payment(db)
    [row] = R._txn_rows(db, [db.get(M.ParkingSession, s.id)])
    assert row['receipt_status'] == 'EXTERNAL' and row['receipt_issuer'] == 'POS терминал (Ontime)'
    ws = load_workbook(BytesIO(asyncio.run(_collect(X.transactions_excel([row]))))).active
    values = [[c.value for c in r] for r in ws.iter_rows()]
    col = values[0].index('Баримтын төлөв')
    assert values[1][col] == 'POS терминал гаргасан (Ontime)'
    assert len(values[0]) == len(values[1])


# ══ 2026-10-14: толгойгүй хуучин апп — POS аппын хэрэглэгч (role / ≥20 карт), веб хөтөч биш ══
CHROME = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/141.0 Safari/537.36'


def _card_history(db, s, u, n):
    for i in range(n):
        db.add(M.Payment(session_id=s.id, site_id=s.site_id, provider='POS', payment_method='CARD',
                         sender_invoice_no=f'hist-card-{u.username}-{i}', amount=Decimal('1000'),
                         vat_amount=Decimal('0'), status='PAID', paid_at=datetime.utcnow(),
                         cashier_id=u.id))
    db.commit()


def test_headerless_fallback_role_and_card_history_but_not_browser(db):
    """Толгой илгээдэггүй хуучин build: кассчин нь role POS эсвэл POS аппаар ≥20 карт авсан бол
    POS апп (зогсоолоос үл хамааран); веб хөтчийн (Mozilla UA) хүсэлт — веб касс → систем баримт;
    картын түүх 19 → POS апп биш."""
    s, u = parking(db), user(db, 'pos.ylalt')
    other = parking(db, code='YLALT1', plate='3333УБА')
    _card_history(db, other, u, 20)                                    # түүх ӨӨР зогсоолд
    out = cash(db, s, u, ua='Dart/3.5 (dart:io)')
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and p.raw_payload['pos_app_via'] == 'card_history'
    assert receipts(db, p.id)[0].status == 'EXTERNAL'
    log = db.query(M.AuditLog).filter_by(action='CASH_PAYMENT', entity_id=p.id).one()
    assert log.detail['pos_app_via'] == 'card_history' and log.detail['user_agent'] == 'Dart/3.5 (dart:io)'
    s2 = another_session(db, s.site_id, '5151УБА')
    out = cash(db, s2, u, ua=CHROME)                                   # тэр хэрэглэгч ВЕБ кассаар
    p2 = db.get(M.Payment, out['payment_id'])
    assert p2.source is None and receipts(db, p2.id)[0].provider == 'MSGBILL' and len(jobs(db, p2.id)) == 1
    # role POS — картын түүхгүй ч
    r = user(db, 'pos.new')
    r.role = 'POS'
    db.commit()
    s3 = another_session(db, s.site_id, '5252УБА')
    p3 = db.get(M.Payment, cash(db, s3, r)['payment_id'])
    assert p3.raw_payload['pos_app_via'] == 'role' and receipts(db, p3.id)[0].status == 'EXTERNAL'
    # 19 карт — босгоос доош
    k = user(db, 'pos.few')
    _card_history(db, other, k, 19)
    s4 = another_session(db, s.site_id, '5353УБА')
    p4 = db.get(M.Payment, cash(db, s4, k)['payment_id'])
    assert p4.source is None and receipts(db, p4.id)[0].provider == 'MSGBILL'


def test_header_wins_even_from_browser_and_company_stays_system(db):
    s, u = parking(db), user(db, 'sogii')
    out = cash(db, s, u, app_version='1.5.3', ua=CHROME)
    assert receipts(db, out['payment_id'])[0].status == 'EXTERNAL'
    s2 = another_session(db, s.site_id, '6161УБА')
    out = asyncio.run(PR.cash_payment({'session_id': s2.id, 'customer_tin': '5317878'}, db, u,
                                      req('1.5.3')))
    [r] = receipts(db, out['payment_id'])
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING')           # ААН — систем гаргана


def test_monnis_site_pos_card_is_external_by_default(db):
    """Моннис (өөр түрээслэгч) — эзний 10-14-ний шийдвэрээр ч мөн терминал гаргана."""
    t = M.Tenant(name='Моннис пропэрти', code='MONNIS')
    db.add(t); db.flush()
    s, u = parking(db, code='MONNIS', plate='7070УБА'), user(db, 'Building')
    db.get(M.ParkingSite, s.site_id).tenant_id = t.id
    db.commit()
    out = card(db, s, u)
    assert receipts(db, out['payment_id'])[0].status == 'EXTERNAL'


def test_file_mode_never_cancels_reissued_receipt():
    p = {'method': 'CARD', 'company': False, 'reissued': True,
         'active': [{'id': 'r', 'provider': 'MSGBILL', 'provider_ref': 'x', 'status': 'SENT',
                     'ebarimt_id': '0' * 25 + '10045952', 'amount': 1000}],
         'primary': {'id': 'r'}, 'amount': 1000}
    assert 'ДАХИН гаргасан' in V.cancel_block_reason(p)
    p['reissued'] = False
    assert V.cancel_block_reason(p) is None


def test_pos_confirm_company_payer_keeps_system_receipt_by_default(db):
    """ААН регистртэй картын төлбөр (POS апп) — терминал иргэний баримт хэвлэдэг тул ерөнхий
    ON үед ч систем ААН-ийн баримтыг msgbill-ээр гаргана (нэг бодлого)."""
    s, u = parking(db), user(db)
    out = asyncio.run(PR.pos_confirm({'session_id': s.id, 'amount': 1000, 'transaction_id': 'bank-co',
                                      'payer_reg_no': '5317878'}, req('1.5.3'), db, u))
    p = db.get(M.Payment, out['payment_id'])
    assert p.source == 'POSAPP' and p.ebarimt_receiver_type == 'COMPANY'
    [r] = receipts(db, p.id)
    assert (r.provider, r.status) == ('MSGBILL', 'PENDING') and len(jobs(db, p.id)) == 1


def test_transfer_and_qr_from_pos_app_are_unchanged(db):
    """Дансаар/QR — POS аппаас ирсэн ч терминал хэвлэдэггүй (карт/бэлэн биш) → систем."""
    s = parking(db)
    for method, provider in (('TRANSFER', 'TRANSFER'), ('QR', 'QPAY')):
        p = M.Payment(session_id=s.id, site_id=s.site_id, provider=provider, payment_method=method,
                      source='POSAPP', sender_invoice_no=f'x-{method}', amount=Decimal('1000'),
                      vat_amount=Decimal('0'), status='PAID', paid_at=datetime.utcnow())
        db.add(p); db.flush()
        assert J.terminal_issues_receipt(db, p) is False
