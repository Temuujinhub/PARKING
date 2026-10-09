"""Synthetic database only; bounded reads, authorization and destructive-action guards."""
from io import BytesIO
from datetime import datetime, timedelta
from hashlib import sha256
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import auth, models as M
from app.database import Base, get_db
from app.routers import admin_router as A, integration_router as I


@pytest.fixture
def ctx():
    engine = create_engine('sqlite://', poolclass=StaticPool, connect_args={'check_same_thread': False})
    Base.metadata.create_all(engine)
    with Session(engine, autoflush=False) as db:
        user = M.User(username='synthetic', password_hash='unused', role='SUPER_ADMIN')
        db.add(user); db.commit()
        app = FastAPI(); app.include_router(A.router); app.include_router(I.router)
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[auth.get_current_user] = lambda: user
        with TestClient(app) as client:
            yield db, user, client, engine
    engine.dispose()


def driver(db, **kw):
    values = dict(plate_number='1234УБА', full_name='Synthetic', company='Company',
                  valid_from=datetime.utcnow()-timedelta(days=1),
                  valid_to=datetime.utcnow()+timedelta(days=30))
    values.update(kw)
    row = M.RegisteredDriver(**values); db.add(row); db.flush(); return row


def test_cursor_complete_order_with_duplicates_and_changed_filter(ctx):
    db, _, c, _ = ctx
    for i in range(207): driver(db, company=['', 'B', 'A'][i % 3], plate_number=f'{i%7:04}УБА')
    db.commit()
    seen=[]; cursor=''
    while True:
        response=c.get('/api/admin/drivers/page', params={'cursor':cursor} if cursor else {})
        assert response.status_code==200, response.text
        data=response.json(); assert data['total']==207 and len(data['items'])<=100
        seen.extend(r['id'] for r in data['items'])
        cursor=data['next_cursor']
        if not cursor: break
        assert c.get('/api/admin/drivers/page',params={'cursor':cursor,'q':'changed'}).status_code==400
    assert len(seen)==len(set(seen))==207
    assert c.get('/api/admin/drivers/page?limit=201').status_code==422
    assert c.get('/api/admin/drivers/page?cursor=bad').status_code==400


def test_page_joins_sites_with_bounded_queries(ctx):
    db, _, c, engine=ctx
    for i in range(105):
        site=M.ParkingSite(site_code=f'S{i}',name=f'Site {i}'); db.add(site); db.flush()
        driver(db,site_id=site.id)
    db.commit(); db.expire_all()
    queries=[]
    def collect(conn,cursor,statement,params,context,many):
        if statement.lstrip().upper().startswith('SELECT'): queries.append(statement)
    event.listen(engine,'before_cursor_execute',collect)
    try: data=c.get('/api/admin/drivers/page').json()
    finally: event.remove(engine,'before_cursor_execute',collect)
    assert len(data['items'])==100 and data['total']==105
    assert len(queries)<=3  # current user, count and a joined page; no per-row site SELECTs


def tenant_setup(db,user):
    own=M.Tenant(name='Own',code='O'); foreign=M.Tenant(name='Foreign',code='F')
    db.add_all([own,foreign]); db.flush()
    user.role='ADMIN'; user.tenant_id=own.id
    first=driver(db,tenant_id=own.id,company='Own')
    other=driver(db,tenant_id=foreign.id,company='Foreign')
    db.commit(); return first,other


def test_tenant_scope_across_page_export_company_and_bulk_delete(ctx):
    db,user,c,_=ctx; first,other=tenant_setup(db,user)
    assert [r['id'] for r in c.get('/api/admin/drivers/page').json()['items']]==[first.id]
    assert c.get('/api/admin/drivers/company-search?q=Foreign').json()==[]
    assert c.get('/api/admin/drivers/options').json()==[]
    for endpoint in ('export','bulk-delete'):
        assert c.post('/api/admin/drivers/'+endpoint,json={'ids':[first.id,other.id]}).status_code==403
    assert db.query(M.RegisteredDriver).count()==2


def test_filtered_export_not_capped_to_page_and_text_is_not_formula(ctx):
    db,_,c,_=ctx
    site=M.ParkingSite(site_code='EXPORT',name='Export'); db.add(site); db.flush()
    for i in range(205): driver(db,site_id=site.id, full_name='=HYPERLINK("https://example.invalid")', company='Export',note='old\x01note')
    driver(db,company='Elsewhere'); db.commit()
    response=c.post('/api/admin/drivers/export',json={'site_id':site.id})
    assert response.status_code==200 and response.headers['cache-control']=='no-store'
    book=load_workbook(BytesIO(response.content)); rows=list(book.active)
    assert len(rows)==206 and rows[1][2].data_type=='s' and rows[1][2].value.startswith('=')
    assert rows[1][5].value=='oldnote'
    ids=[rows[1][0].value, rows[2][0].value]
    selected=c.post('/api/admin/drivers/export',json={'ids':ids})
    assert len(list(load_workbook(BytesIO(selected.content)).active))==3
    assert c.post('/api/admin/drivers/export',json={'ids':[]}).status_code==422


def test_bulk_delete_preview_stale_and_atomic_audit(ctx,monkeypatch):
    db,_,c,_=ctx; rows=[driver(db),driver(db)]; db.commit()
    ids=[r.id for r in rows]
    body={'ids':ids}
    proof=c.post('/api/admin/drivers/bulk-delete',json=body).json()
    assert proof['selected']==2 and db.query(M.RegisteredDriver).count()==2
    assert db.query(M.AuditLog).count()==0
    rows[0].note='Changed'; db.commit()
    assert c.post('/api/admin/drivers/bulk-delete',json={**body,'dry_run':False,'preview_token':proof['preview_token']}).status_code==409
    proof=c.post('/api/admin/drivers/bulk-delete',json=body).json()
    response=c.post('/api/admin/drivers/bulk-delete',json={**body,'dry_run':False,'preview_token':proof['preview_token']})
    assert response.status_code==200 and response.json()['deleted']==2
    assert db.query(M.RegisteredDriver).count()==0
    audit=db.query(M.AuditLog).filter_by(action='DELETE').all()
    assert len(audit)==2 and audit[0].detail['before']['plate_number']=='1234УБА'
    assert db.query(M.Payment).count()==db.query(M.BarrierCommand).count()==0


@pytest.mark.parametrize('role,permissions',[('OPERATOR',['drivers']),('ADMIN',[])])
def test_delete_permissions(ctx,role,permissions):
    db,user,c,_=ctx; row=driver(db); user.role=role;user.permissions=permissions; db.commit()
    assert c.post('/api/admin/drivers/bulk-delete',json={'ids':[row.id]}).status_code==403


@pytest.mark.parametrize('payload',[
    {'latitude':91,'longitude':10}, {'latitude':47}, {'longitude':106},
    {'google_maps_url':'https://google.com.evil.example/maps/abc'},
    {'google_maps_url':'http://maps.google.com/?q=1,2'},
    {'google_maps_url':'https://user:pass@maps.google.com/?q=1,2'},
    {'google_maps_url':'https://maps.google.com:8443/?q=1,2'},
])
def test_invalid_location_rejected(ctx,payload):
    db,_,c,_=ctx
    response=c.post('/api/admin/sites',json={'name':'Synthetic','site_code':'LOCATION',**payload})
    assert response.status_code==422, response.text
    assert db.query(M.ParkingSite).count()==0


def test_map_location_roundtrip_and_clear_pair(ctx):
    db,_,c,_=ctx
    response=c.post('/api/admin/sites',json={'name':'Synthetic','site_code':'LOCATION',
        'google_maps_url':'https://maps.app.goo.gl/abc','latitude':47.918,'longitude':106.917})
    assert response.status_code==200,response.text
    sid=response.json()['id']
    assert c.put('/api/admin/sites/'+sid,json={'latitude':None}).status_code==422
    assert c.put('/api/admin/sites/'+sid,json={'name':'Renamed'}).status_code==200
    assert float(db.get(M.ParkingSite,sid).latitude)==47.918
    assert c.put('/api/admin/sites/'+sid,json={'latitude':None,'longitude':None}).status_code==200


def test_partner_map_auth_scope_nested_counts_and_missing_capacity(ctx):
    db,_,c,_=ctx
    parent=M.ParkingSite(name='Parent',site_code='P',capacity=2,latitude=47.9,longitude=106.9)
    single=M.ParkingSite(name='Single',site_code='S',capacity=0)
    hidden=M.ParkingSite(name='Inactive',site_code='H',is_active=False)
    db.add_all([parent,single,hidden]);db.flush()
    child=M.ParkingSite(name='Child',site_code='C',parent_site_id=parent.id,capacity=1)
    db.add(child); db.flush()
    for i in range(4): db.add(M.ParkingSession(site_id=parent.id,plate_number=f'P{i}',entry_time=datetime.utcnow(),status='OPEN',paused_since=datetime.utcnow() if i==0 else None))
    db.add(M.ParkingSession(site_id=single.id,plate_number='S1',entry_time=datetime.utcnow(),status='OPEN',paused_since=datetime.utcnow()))
    db.add(M.ParkingSession(site_id=parent.id,plate_number='CLOSED',entry_time=datetime.utcnow(),status='CLOSED'))
    key=M.PartnerKey(name='MAP',key_hash=sha256(b'synthetic-map-key').hexdigest(),key_prefix='synthetic',scopes='read',site_id=parent.id)
    db.add(key);db.commit()
    assert c.get('/api/v1/sites').status_code==401
    response=c.get('/api/v1/sites',headers={'X-API-Key':'synthetic-map-key'})
    assert response.status_code==200,response.text
    row=response.json()['sites'][0]
    assert len(response.json()['sites'])==1
    assert row['occupied']==3 and row['free']==0 and row['occupancy_percent']==150
    assert row['over_capacity'] and row['location_available'] and row['occupancy_is_estimate']
    assert 'plate_number' not in row and 'qpay_password' not in row
    admin=next(s for s in c.get('/api/admin/sites').json() if s['id']==parent.id)
    assert admin['occupied']==row['occupied']
    key.site_id=single.id;db.commit()
    row=c.get('/api/v1/sites',headers={'X-API-Key':'synthetic-map-key'}).json()['sites'][0]
    assert row['occupied']==1 and row['free'] is None and row['occupancy_percent'] is None
    assert not row['location_available'] and row['latitude'] is None


def test_select_filtered_crosses_pages_and_rejects_over_limit(ctx):
    db,_,c,_=ctx
    db.add_all([M.RegisteredDriver(plate_number=f'{i:04}УБА',company='Group',is_active=False,
                valid_to=datetime(2027,1,1)) for i in range(2001)])
    db.commit()
    assert c.post('/api/admin/drivers/select-filtered',json={}).status_code==413
    row=db.query(M.RegisteredDriver).first();row.is_active=True;db.commit()
    selected=c.post('/api/admin/drivers/select-filtered',json={'is_active':False}).json()
    assert selected['selected']==len(set(selected['ids']))==2000 and row.id not in selected['ids']
    assert c.get('/api/admin/drivers/page?site_id=invalid').status_code==422


def test_delete_audit_failure_rolls_back_all_rows(ctx,monkeypatch):
    from app.routers.driver_directory_router import delete_drivers, DeleteRequest
    db,user,_,_=ctx
    rows=[driver(db),driver(db)];db.commit()
    body=DeleteRequest(ids=[r.id for r in rows])
    proof=delete_drivers(body,db,user)
    original=A._audit
    def fail_summary(db,user,action,*args,**kw):
        if action=='BULK_DELETE': raise RuntimeError('synthetic audit failure')
        return original(db,user,action,*args,**kw)
    monkeypatch.setattr(A,'_audit',fail_summary)
    with pytest.raises(RuntimeError,match='audit failure'):
        delete_drivers(body.model_copy(update={'dry_run':False,'preview_token':proof['preview_token']}),db,user)
    db.rollback()  # get_db closes/rolls back failed request sessions in the application
    assert db.query(M.RegisteredDriver).count()==2 and db.query(M.AuditLog).count()==0


def test_empty_site_scope_cannot_leak_global_rows(ctx,monkeypatch):
    from app.services import driver_directory
    db,user,c,_=ctx
    driver(db);db.commit()
    monkeypatch.setattr(driver_directory,'operator_sites',lambda _: [])
    assert c.get('/api/admin/drivers/page').json()['items']==[]
    assert c.post('/api/admin/drivers/select-filtered',json={}).json()['ids']==[]
    book=load_workbook(BytesIO(c.post('/api/admin/drivers/export',json={}).content))
    assert len(list(book.active))==1


def test_tenant_cannot_edit_other_site_location(ctx):
    db,user,c,_=ctx
    tenant_setup(db,user)
    site=M.ParkingSite(name='Foreign site',site_code='OTHER')
    db.add(site);db.commit()
    assert c.put('/api/admin/sites/'+site.id,json={'latitude':47.9,'longitude':106.9}).status_code==403
    assert db.get(M.ParkingSite,site.id).latitude is None
