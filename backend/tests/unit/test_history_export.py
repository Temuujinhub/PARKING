"""Түүх: УБ-ын өдрийн огнооны шүүлт (№35), гарсан цагаар эрэмбэ/шүүлт (№59),
байгууллагаар шүүх (№61), арын Excel экспорт (№62) — system test 2026-10-22."""
import json
from datetime import datetime
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app import database
from app import models as M
from app.config import settings
from app.database import Base
from app.routers import sessions_router as R
from app.services import app_settings as A
from app.services import history_export as HX

ADMIN = M.User(id='abcdef11-1111-4111-8111-11111111abcd', username='admin', role='SUPER_ADMIN', is_active=True)


@pytest.fixture
def engine():
    eng = create_engine('sqlite://', poolclass=StaticPool)
    Base.metadata.create_all(eng)
    A.invalidate_cache()
    yield eng
    A.invalidate_cache()
    eng.dispose()


@pytest.fixture
def db(engine):
    with Session(engine, autoflush=False) as session:
        yield session


@pytest.fixture
def site(db):
    tariff = M.TariffTemplate(name='Hourly', free_minutes=0, extra_hour_price=1000,
                              daily_cap=50000, grace_minutes=15)
    s = M.ParkingSite(name='Хангарьд', site_code='HANG', tariff_template=tariff,
                      auto_close_hours=0, entry_only_free_hours=0)
    db.add(s); db.flush()
    return s


def stay(db, site, plate, entry, exit=None, status='CLOSED', **kw):
    s = M.ParkingSession(site_id=site.id, plate_number=plate, entry_time=entry, exit_time=exit,
                         status=status, total_fee=kw.pop('total_fee', 1000), **kw)
    db.add(s); db.flush()
    return s


def plates(result):
    return [r['plate_number'] for r in result['rows']]


def test_date_filter_uses_ulaanbaatar_day_not_utc(db, site):
    # УБ 09-03 01:00 (UTC 09-02 17:00) — 09-03-нд ОРНО; өмнө нь алга болдог байв
    stay(db, site, '0001УБА', datetime(2026, 9, 2, 17, 0))
    stay(db, site, '0002УБА', datetime(2026, 9, 3, 1, 0))     # УБ 09-03 09:00
    # УБ 09-04 01:00 (UTC 09-03 17:00) — 09-03-нд ОРОХГҮЙ; өмнө нь орж ирдэг байв
    stay(db, site, '0003УБА', datetime(2026, 9, 3, 17, 0))
    db.commit()
    got = R.list_sessions(db=db, user=ADMIN, date_from='2026-09-03', date_to='2026-09-03')
    assert sorted(plates(got)) == ['0001УБА', '0002УБА']
    assert R._local_day_bounds('2026-09-03', '2026-09-03') == (
        datetime(2026, 9, 2, 16, 0), datetime(2026, 9, 3, 16, 0))


def test_bad_date_is_400_not_500(db, site):
    with pytest.raises(HTTPException) as e:
        R.list_sessions(db=db, user=ADMIN, date_from='03/09/2026')
    assert e.value.status_code == 400


def test_exit_date_field_and_exit_sort_keep_open_stays_last(db, site):
    stay(db, site, '0001УБА', datetime(2026, 9, 3, 1), datetime(2026, 9, 3, 9))
    stay(db, site, '0002УБА', datetime(2026, 9, 3, 2), datetime(2026, 9, 3, 3))
    stay(db, site, '0003УБА', datetime(2026, 9, 3, 4), None, status='OPEN')
    stay(db, site, '0004УБА', datetime(2026, 9, 1, 4), datetime(2026, 9, 3, 5))  # өмнө орсон
    db.commit()
    by_exit = R.list_sessions(db=db, user=ADMIN, date_from='2026-09-03', date_to='2026-09-03',
                              date_field='exit', sort='exit_asc')
    assert plates(by_exit) == ['0002УБА', '0004УБА', '0001УБА']     # гараагүй нь шүүлтэд орохгүй
    all_rows = R.list_sessions(db=db, user=ADMIN, sort='exit_desc')
    assert plates(all_rows)[-1] == '0003УБА'                         # NULL гарсан цаг төгсгөлд
    assert plates(all_rows)[0] == '0001УБА'
    with pytest.raises(HTTPException):
        R.list_sessions(db=db, user=ADMIN, sort='plate')
    with pytest.raises(HTTPException):
        R.list_sessions(db=db, user=ADMIN, date_field='paid')


def test_company_filter_matches_registered_plates_case_insensitive(db, site):
    # Кирилл үсгийн том/жижгийг PostgreSQL lower() (C.UTF-8) тэгшитгэнэ; SQLite
    # зөвхөн ASCII-г тэгшитгэдэг тул энд латин нэрээр шалгана.
    db.add(M.RegisteredDriver(plate_number='0001УБА', site_id=site.id, company='Kids Palace',
                              contract_type='CONTRACT', valid_to=datetime(2027, 1, 1)))
    db.flush()
    stay(db, site, '0001УБА', datetime(2026, 9, 3, 1))
    stay(db, site, '0002УБА', datetime(2026, 9, 3, 2))
    db.commit()
    got = R.list_sessions(db=db, user=ADMIN, company='  kids PALACE ')
    assert plates(got) == ['0001УБА']


def test_sync_excel_has_payment_wait_column(db, site):
    import asyncio
    s = stay(db, site, '0001УБА', datetime(2026, 9, 3, 1), status='AWAITING_PAYMENT',
             payment_wait_started_at=datetime(2026, 9, 3, 2, 30))
    db.commit()
    response = R.sessions_excel(db=db, user=ADMIN, plate=s.plate_number)

    async def read():
        return b''.join([part async for part in response.body_iterator])
    rows = list(load_workbook(BytesIO(asyncio.run(read())), read_only=True).active.values)
    assert rows[0][-1] == 'Төлбөр хүлээж эхэлсэн'
    assert rows[1][-1] == '2026-09-03 10:30'          # УБ цаг


@pytest.fixture
def export_env(engine, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'history_export_dir', str(tmp_path))
    monkeypatch.setattr(database, 'SessionLocal', sessionmaker(bind=engine, autoflush=False))
    return tmp_path


def test_background_export_writes_all_rows_beyond_old_20000_cap(db, site, export_env, monkeypatch):
    monkeypatch.setattr(HX, 'CHUNK', 3)                 # олон chunk-ээр урсгахыг шалгана
    owner = M.User(id=ADMIN.id, username='admin', role='SUPER_ADMIN', password_hash='x')
    db.add(owner)
    for i in range(7):
        stay(db, site, f'{i:04d}УБА', datetime(2026, 9, 3, 1, i), datetime(2026, 9, 3, 2, i))
    db.commit()
    job = {'id': 'a' * 32, 'status': 'queued', 'user_id': ADMIN.id, 'username': 'admin',
           'filters': {'site_id': None, 'status': None, 'plate': None, 'date_from': '2026-09-03',
                       'date_to': '2026-09-03', 'inner': None, 'debt': 0, 'date_field': 'entry',
                       'company': None},
           'sort': 'exit_asc', 'filename': 'x.xlsx', 'max_rows': 5}
    HX._write_job(job)
    assert HX.run(str(export_env / f"{job['id']}.json")) == 0
    done = HX.read_job(job['id'])
    assert done['status'] == 'done' and done['done'] == 5 and done['total'] == 5
    assert done['truncated'] is True                    # 7 мөр > max_rows 5 → анхааруулна
    rows = list(load_workbook(export_env / f"{job['id']}.xlsx", read_only=True).active.values)
    assert rows[0][0] == 'Дугаар' and len(rows[0]) == len(HX.HEADERS)
    assert [r[0] for r in rows[1:6]] == ['0000УБА', '0001УБА', '0002УБА', '0003УБА', '0004УБА']
    assert rows[6][0] == 'Нийт' and rows[6][7] == 5000
    assert rows[7][0].startswith('⚠ 7 мөрөөс')


def test_failed_export_records_error_and_leaves_no_partial_file(db, site, export_env):
    job = {'id': 'b' * 32, 'status': 'queued', 'user_id': 'missing', 'filters': {}, 'sort': None}
    HX._write_job(job)
    assert HX.run(str(export_env / f"{job['id']}.json")) == 1
    assert HX.read_job(job['id'])['status'] == 'failed'
    assert not list(export_env.glob('*.xlsx*'))


def test_dead_running_job_becomes_failed(export_env):
    HX._write_job({'id': 'c' * 32, 'status': 'running', 'pid': 999999999})
    assert HX.refresh(HX.read_job('c' * 32))['status'] == 'failed'


def test_start_export_validates_range_and_single_slot(db, site, export_env, monkeypatch):
    started = []
    monkeypatch.setattr(HX, 'start_job', lambda user, f, sort, name: started.append((f, sort, name))
                        or {'id': 'd' * 32, 'status': 'queued'})
    with pytest.raises(HTTPException) as e:
        R.start_export({}, db=db, user=ADMIN)                       # огноо ч, дугаар ч алга
    assert e.value.status_code == 400
    with pytest.raises(HTTPException) as e:
        R.start_export({'date_from': '2026-06-01', 'date_to': '2026-09-30'}, db=db, user=ADMIN)
    assert e.value.status_code == 400                               # 122 хоног > 93
    with pytest.raises(HTTPException):
        R.start_export({'date_from': '2026-09-01', 'date_to': '2026-09-30', 'sort': 'x'}, db=db, user=ADMIN)
    out = R.start_export({'date_from': '2026-07-01', 'date_to': '2026-09-30', 'debt': '1',
                          'date_field': 'exit', 'sort': 'exit_desc'}, db=db, user=ADMIN)
    assert out['status'] == 'queued'
    f, sort, name = started[0]
    assert f['debt'] == 1 and f['date_field'] == 'exit' and sort == 'exit_desc'
    assert name == 'tuuh_2026-07-01_2026-09-30.xlsx'
    # Нэг дугаарын бүх түүх — огноогүй ч болно
    R.start_export({'plate': '1234УБА'}, db=db, user=ADMIN)
    # Өөр экспорт явагдаж байвал 409
    monkeypatch.setattr(HX, 'start_job', lambda *a: (_ for _ in ()).throw(
        RuntimeError({'username': 'tumee', 'done': 10, 'total': 100})))
    with pytest.raises(HTTPException) as e:
        R.start_export({'plate': '1234УБА'}, db=db, user=ADMIN)
    assert e.value.status_code == 409


def test_job_files_are_private_to_owner(export_env):
    HX._write_job({'id': 'e' * 32, 'status': 'done', 'user_id': 'someone-else'})
    other = M.User(id='abcdef22-2222-4222-8222-22222222abcd', username='op', role='OPERATOR')
    with pytest.raises(HTTPException) as e:
        R.export_status('e' * 32, user=other)
    assert e.value.status_code == 404
    with pytest.raises(HTTPException):
        R.export_status('../../etc/passwd', user=ADMIN)
    assert R.export_status('e' * 32, user=ADMIN)['status'] == 'done'   # SUPER_ADMIN харна
    assert json.loads((export_env / ('e' * 32 + '.json')).read_text())['user_id'] == 'someone-else'
